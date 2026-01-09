from __future__ import annotations

import re
from datetime import datetime
from pathlib import Path

from docx import Document
from job_mcp.exceptions import ValidationError


def _safe_filename(name: str) -> str:
    """Make a safe filename for Windows/Linux/macOS ."""
    name = (name or "").strip() or "resume"
    name = re.sub(r'[<>:"/\\|?*\n\r\t]+', " ", name)
    name = re.sub(r"\s+", " ", name).strip()
    return name[:120]


def export_resume_to_docx(
    resume_text: str,
    output_dir: str = "output_resumes",
    file_name: str | None = None,
    add_timestamp: bool = True,
) -> str:
    """
    Export resume plain text to a DOCX file.
    
    Creates a professional DOCX file with proper paragraph formatting,
    suitable for ATS (Applicant Tracking System) scanning.
    
    Args:
        resume_text: The resume content (plain text).
        output_dir: Directory to save the file (relative or absolute).
        file_name: Base filename without extension (optional).
        add_timestamp: Whether to append YYYYMMDD_HHMMSS to filename.
    
    Returns:
      absolute path to the generated docx file (string)
    """
    text = (resume_text or "").strip()
    if not text:
        raise ValidationError("resume_text cannot be empty")

    out_dir = Path(output_dir).expanduser()
    if not out_dir.is_absolute():
        out_dir = (Path.cwd() / out_dir).resolve()
    out_dir.mkdir(parents=True, exist_ok=True)

    base = _safe_filename(file_name) if file_name else "rewritten_resume"
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S") if add_timestamp else ""
    final_name = f"{base}_{stamp}.docx" if stamp else f"{base}.docx"
    out_path = (out_dir / final_name).resolve()

    # Create DOCX document
    doc = Document()

    # Convert text blocks into paragraphs (simple + ATS-friendly)
    blocks = text.split("\n")
    for line in blocks:
        line = line.rstrip()
        if line == "":
            doc.add_paragraph("")  # keep spacing
        else:
            doc.add_paragraph(line)

    doc.save(str(out_path))
    return str(out_path)
