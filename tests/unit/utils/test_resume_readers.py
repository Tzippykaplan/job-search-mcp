import pytest
from docx import Document
from pathlib import Path

from job_mcp.utils.read_resume import read_text_file, read_docx_file, read_resume_any
from job_mcp.exceptions import FileReadError


def test_read_text_file_reads_utf8(tmp_path):
    p = tmp_path / "resume.txt"
    p.write_text("Hello\nWorld", encoding="utf-8")
    assert read_text_file(str(p)) == "Hello\nWorld"


def test_read_docx_file_joins_non_empty_paragraphs(tmp_path):
    p = tmp_path / "resume.docx"

    doc = Document()
    doc.add_paragraph("  First  ")
    doc.add_paragraph("")          # empty
    doc.add_paragraph("   ")       # whitespace only
    doc.add_paragraph("Second")
    doc.save(str(p))

    assert read_docx_file(str(p)) == "First\nSecond"


def test_read_resume_any_txt_routes_to_text(allow_tmp_dir):
    p = allow_tmp_dir / "resume.txt"
    p.write_text("Text resume", encoding="utf-8")
    assert read_resume_any(str(p)) == "Text resume"


def test_read_resume_any_md_routes_to_text(allow_tmp_dir):
    p = allow_tmp_dir / "resume.md"
    p.write_text("# Title\nBody", encoding="utf-8")
    assert read_resume_any(str(p)) == "# Title\nBody"


def test_read_resume_any_docx_routes_to_docx(allow_tmp_dir):
    p = allow_tmp_dir / "resume.docx"
    doc = Document()
    doc.add_paragraph("Line1")
    doc.add_paragraph("Line2")
    doc.save(str(p))

    assert read_resume_any(str(p)) == "Line1\nLine2"


def test_read_resume_any_missing_file_raises(allow_tmp_dir):
    p = allow_tmp_dir / "missing.txt"
    with pytest.raises(FileReadError):
        read_resume_any(str(p))


def test_read_resume_any_unsupported_extension_raises(allow_tmp_dir):
    p = allow_tmp_dir / "resume.pdf"
    p.write_text("nope", encoding="utf-8")
    with pytest.raises(FileReadError, match="Unsupported resume file type"):
        read_resume_any(str(p))
